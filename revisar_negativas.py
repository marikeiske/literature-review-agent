"""Planilha de revisão humana das perguntas negativas (HITL sobre a avaliação).

A revisão se faz na interface (python app.py, aba Revisar), que grava em
resultados/revisao_<rodada>.csv. Esta planilha é a versão para anexar ou mostrar: põe lado
a lado a resposta esperada, a resposta do agente, a nota automática e a sua revisão.

Uso:
  python revisar_negativas.py                         todos os resultados/neg_*.csv
  python revisar_negativas.py resultados/neg_X.csv    só um

Saída: resultados/revisao_<nome do csv>.xlsx, com o que você marcou na interface.
"""
import csv
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import revisao
from negativas import NEGATIVAS

RAIZ = Path(__file__).resolve().parent
ESPERADO = {n["id"]: n["esperado"] for n in NEGATIVAS}
CABECALHO = ["ID", "Tipo", "Pergunta", "Resposta esperada", "Resposta do agente",
             "Automático", "Sua revisão", "Observação"]
LARGURAS = [7, 16, 38, 45, 70, 22, 14, 35]


def automatico(l: dict) -> str:
    marcas = [m for m, c in (("recusou", "recusou"), ("corrigiu", "corrigiu"), ("INVENTOU", "inventou"),
                             ("recusou sem ler", "recusa_sem_ler")) if l.get(c) == "1"]
    return ("ACERTO" if l["acerto"] == "1" else "erro") + (" · " + ", ".join(marcas) if marcas else "")


def gerar(csv_path: Path) -> Path:
    saida = csv_path.parent / f"revisao_{csv_path.stem}.xlsx"
    anteriores = {}
    if saida.exists():
        ws_old = openpyxl.load_workbook(saida)["Revisão"]
        for r in ws_old.iter_rows(min_row=2, values_only=True):
            if r[0]:
                anteriores[r[0]] = (r[6], r[7])
    # O que você marcou na tela Revisar da interface tem prioridade sobre a planilha
    for id_, rev in revisao.ler_revisoes(csv_path.parent / f"revisao_{csv_path.stem}.csv").items():
        if rev.get("status") or rev.get("comentario"):
            anteriores[id_] = (rev.get("status") or None, rev.get("comentario") or None)

    linhas = {}
    with open(csv_path, encoding="utf-8-sig") as f:
        for l in csv.DictReader(f, delimiter=";"):
            linhas.setdefault(l["id"], l)  # id repetido: vale a primeira execução

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Revisão"
    ws.append(CABECALHO)
    for l in sorted(linhas.values(), key=lambda x: x["id"]):
        marca, obs = anteriores.get(l["id"], (None, None))
        ws.append([l["id"], l["tipo"], l["pergunta"], ESPERADO.get(l["id"], ""), l["resposta"],
                   automatico(l), marca, obs])
    n = ws.max_row

    cab = PatternFill("solid", fgColor="1B2A3A")
    sua = PatternFill("solid", fgColor="FFF6CC")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = cab
    for i, w in enumerate(LARGURAS, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        row[6].fill = sua
        row[7].fill = sua
    ws.freeze_panes = "C2"
    dv = DataValidation(type="list", formula1='"Confere,Parcial,Não confere"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G2:G{n}")

    r = wb.create_sheet("Resumo")
    r.column_dimensions["A"].width = 48
    r.column_dimensions["B"].width = 12
    dados = [
        ("Perguntas", f"=COUNTA(Revisão!A2:A{n})"),
        ("Acerto automático (regex)", f'=COUNTIF(Revisão!F2:F{n},"ACERTO*")'),
        ("Revisadas por você", f"=COUNTA(Revisão!G2:G{n})"),
        ("Confere", f'=COUNTIF(Revisão!G2:G{n},"Confere")'),
        ("Parcial", f'=COUNTIF(Revisão!G2:G{n},"Parcial")'),
        ("Não confere", f'=COUNTIF(Revisão!G2:G{n},"Não confere")'),
        ("Divergência: automático ACERTO, você Não confere",
         f'=COUNTIFS(Revisão!F2:F{n},"ACERTO*",Revisão!G2:G{n},"Não confere")'),
        ("Divergência: automático erro, você Confere",
         f'=COUNTIFS(Revisão!F2:F{n},"erro*",Revisão!G2:G{n},"Confere")'),
    ]
    r.append(["Indicador", "Valor"])
    for c in r[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = cab
    for d in dados:
        r.append(list(d))
    r.append([])
    r.append(["A pontuação automática é necessária, não suficiente: o artigo só reporta "
              "como certo o que você marcou como Confere."])
    r.cell(r.max_row, 1).alignment = Alignment(wrap_text=True)

    wb.save(saida)
    return saida


def main():
    arquivos = [Path(a) for a in sys.argv[1:]] or [
        p for p in sorted((RAIZ / "resultados").glob("neg_*.csv")) if not p.stem.endswith("_bruto")]
    if not arquivos:
        sys.exit("Nenhum resultados/neg_*.csv. Rode antes: python avaliar.py <modelo> --negativas")
    for a in arquivos:
        print(f"{a.name} -> {gerar(a).name}")


if __name__ == "__main__":
    main()
